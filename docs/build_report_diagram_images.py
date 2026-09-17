from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "diagram_images"
DOCX_OUT = ROOT / "docs" / "report_diagrams_with_images.docx"


COLORS = {
    "bg": "#FFFFFF",
    "ink": "#17324D",
    "muted": "#52616F",
    "blue": "#D9EAF7",
    "blue_border": "#2E74B5",
    "green": "#E3F4E8",
    "green_border": "#2E7D32",
    "orange": "#FFF1D6",
    "orange_border": "#C77700",
    "red": "#FCE2E2",
    "red_border": "#B3261E",
    "gray": "#F3F5F7",
    "gray_border": "#AAB4BE",
    "purple": "#EFE7FA",
    "purple_border": "#6F42C1",
}


def font(size=26, bold=False):
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


TITLE = font(34, True)
SUBTITLE = font(24, True)
BODY = font(22)
SMALL = font(18)
TINY = font(15)


def wrap(text, width=22):
    return "\n".join(textwrap.wrap(text, width=width))


def text_size(draw, text, fnt):
    box = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=5)
    return box[2] - box[0], box[3] - box[1]


def draw_title(draw, title, width):
    draw.text((width // 2, 34), title, fill=COLORS["ink"], font=TITLE, anchor="ma")
    draw.line((80, 86, width - 80, 86), fill=COLORS["blue_border"], width=3)


def box(draw, xy, text, fill="blue", outline=None, fnt=BODY, radius=18, wrap_width=22):
    x1, y1, x2, y2 = xy
    outline = outline or COLORS[f"{fill}_border"]
    draw.rounded_rectangle(xy, radius=radius, fill=COLORS[fill], outline=outline, width=3)
    wrapped = wrap(text, wrap_width)
    tw, th = text_size(draw, wrapped, fnt)
    draw.multiline_text(((x1 + x2) / 2, (y1 + y2 - th) / 2), wrapped, fill=COLORS["ink"], font=fnt, anchor="ma", align="center", spacing=5)


def arrow(draw, start, end, color="#51606D", width=4):
    draw.line((start, end), fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - 16 * direction, y2 - 9), (x2 - 16 * direction, y2 + 9)]
    else:
        direction = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 9, y2 - 16 * direction), (x2 + 9, y2 - 16 * direction)]
    draw.polygon(pts, fill=color)


def canvas(title, width=1800, height=1100):
    img = Image.new("RGB", (width, height), COLORS["bg"])
    draw = ImageDraw.Draw(img)
    draw_title(draw, title, width)
    return img, draw


def save(img, name):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{name}.png"
    img.save(path, quality=95)
    return path


def architecture():
    img, d = canvas("High-Level System Architecture", 1800, 1200)
    layers = [
        ("Client Layer", ["Web Browser", "Mobile App", "Role-based Users"], 140, "blue"),
        ("React + Vite Frontend", ["React Router", "Axios API Client", "Maps / Charts"], 370, "green"),
        ("Django Backend", ["REST API", "JWT Auth", "Domain Apps", "Channels", "Celery"], 600, "orange"),
        ("Data and Integrations", ["PostgreSQL", "Redis", "Twilio", "AI Triage", "Maps"], 850, "purple"),
    ]
    centers = []
    for title, items, y, color in layers:
        box(d, (90, y, 1710, y + 150), title, fill=color, fnt=SUBTITLE, wrap_width=40)
        item_w = 260
        gap = 35
        total = len(items) * item_w + (len(items) - 1) * gap
        x = (1800 - total) // 2
        for item in items:
            box(d, (x, y + 75, x + item_w, y + 135), item, fill="gray", fnt=SMALL, radius=12, wrap_width=18)
            x += item_w + gap
        centers.append((900, y + 150))
    for idx in range(len(centers) - 1):
        arrow(d, (900, centers[idx][1] + 6), (900, layers[idx + 1][2] - 8))
    return save(img, "01_high_level_system_architecture")


def use_case():
    img, d = canvas("Use Case Diagram", 1900, 1300)
    actors = [
        ("Patient", 110, 220), ("Reception Staff", 110, 460), ("Doctor", 110, 700),
        ("Driver", 1680, 220), ("Supervisor", 1680, 460), ("Admin", 1680, 700),
    ]
    use_cases = [
        ("Register / Sign In", 690, 150), ("Search Hospitals", 690, 270), ("Emergency Triage", 690, 390),
        ("Book Ambulance", 690, 510), ("Manage Beds", 690, 630), ("Transfer Patient", 690, 750),
        ("Monitor Alerts", 690, 870), ("Verify Hospital", 690, 990), ("View Analytics", 690, 1110),
    ]
    d.rounded_rectangle((520, 115, 1380, 1205), radius=28, outline=COLORS["blue_border"], width=4)
    d.text((950, 125), "Healthcare Management System", fill=COLORS["ink"], font=SUBTITLE, anchor="ma")
    for name, x, y in actors:
        d.ellipse((x, y, x + 90, y + 90), fill=COLORS["green"], outline=COLORS["green_border"], width=3)
        d.text((x + 45, y + 112), name, fill=COLORS["ink"], font=SMALL, anchor="ma")
    for name, x, y in use_cases:
        d.ellipse((x, y, x + 520, y + 75), fill=COLORS["blue"], outline=COLORS["blue_border"], width=3)
        d.text((x + 260, y + 38), name, fill=COLORS["ink"], font=BODY, anchor="mm")
    links = [(155, 265, 690, 188), (155, 265, 690, 428), (155, 505, 690, 668), (155, 505, 690, 788),
             (1725, 265, 1210, 548), (1725, 505, 1210, 908), (1725, 505, 1210, 1028), (1725, 745, 1210, 1028), (1725, 745, 1210, 1148)]
    for x1, y1, x2, y2 in links:
        d.line((x1, y1, x2, y2), fill="#6B7785", width=3)
    return save(img, "02_use_case_diagram")


def er_diagram():
    img, d = canvas("Main Database E-R Diagram", 2200, 1500)
    entities = {
        "USERS": (80, 160), "HOSPITALS": (800, 160), "DEPARTMENTS": (1520, 160),
        "PATIENTS": (80, 500), "BEDS": (800, 500), "DOCTORS": (1520, 500),
        "BED_ALLOCATIONS": (80, 840), "AMBULANCES": (800, 840), "EMERGENCY_CASES": (1520, 840),
        "TRANSFER_REQUESTS": (80, 1180), "ALERTS": (800, 1180), "RESOURCE_REQUESTS": (1520, 1180),
    }
    for name, (x, y) in entities.items():
        box(d, (x, y, x + 500, y + 120), name, fill="blue", fnt=SUBTITLE, wrap_width=22)
    relations = [
        ("USERS", "HOSPITALS"), ("HOSPITALS", "DEPARTMENTS"), ("HOSPITALS", "BEDS"),
        ("HOSPITALS", "DOCTORS"), ("USERS", "PATIENTS"), ("PATIENTS", "BED_ALLOCATIONS"),
        ("BEDS", "BED_ALLOCATIONS"), ("HOSPITALS", "AMBULANCES"), ("AMBULANCES", "TRANSFER_REQUESTS"),
        ("HOSPITALS", "EMERGENCY_CASES"), ("HOSPITALS", "ALERTS"), ("BEDS", "ALERTS"),
        ("HOSPITALS", "RESOURCE_REQUESTS"), ("PATIENTS", "TRANSFER_REQUESTS"),
    ]
    def center(name):
        x, y = entities[name]
        return x + 250, y + 60
    for a, b in relations:
        d.line((center(a), center(b)), fill="#59636E", width=3)
    d.text((1100, 1410), "Crow-foot notation simplified: each line represents a primary relationship between tables.", fill=COLORS["muted"], font=SMALL, anchor="ma")
    return save(img, "03_main_database_er_diagram")


def class_diagram():
    img, d = canvas("UML Class Diagram - Core Domain", 2200, 1400)
    classes = [
        ("User", ["id", "email", "role", "hospital"], 90, 160),
        ("Hospital", ["name", "city", "total_beds", "verification_status"], 620, 160),
        ("Department", ["name", "dept_type", "floor"], 1150, 160),
        ("Bed", ["bed_number", "bed_type", "status"], 1680, 160),
        ("Patient", ["full_name", "phone", "blood_group"], 90, 560),
        ("BedAllocation", ["admitted_at", "discharged_at", "is_active"], 620, 560),
        ("Ambulance", ["vehicle_number", "type", "status", "location"], 1150, 560),
        ("AmbulanceRequest", ["pickup", "status", "requested_at"], 1680, 560),
        ("EmergencyCase", ["case_id", "ai_p_level", "needs_profile", "status"], 620, 960),
    ]
    for name, attrs, x, y in classes:
        d.rounded_rectangle((x, y, x + 430, y + 260), radius=12, fill=COLORS["gray"], outline=COLORS["gray_border"], width=3)
        d.rectangle((x, y, x + 430, y + 60), fill=COLORS["blue"], outline=COLORS["blue_border"], width=3)
        d.text((x + 215, y + 31), name, fill=COLORS["ink"], font=SUBTITLE, anchor="mm")
        for idx, attr in enumerate(attrs):
            d.text((x + 26, y + 90 + idx * 38), f"+ {attr}", fill=COLORS["ink"], font=BODY)
    links = [((835, 290), (620, 290)), ((1050, 290), (1150, 290)), ((1580, 290), (1680, 290)),
             ((305, 690), (620, 690)), ((835, 560), (835, 420)), ((1365, 690), (1680, 690)),
             ((835, 960), (835, 420))]
    for s, e in links:
        arrow(d, s, e)
    return save(img, "04_uml_class_diagram_core_domain")


def vertical_flow(title, steps, name, width=1600):
    height = 180 + len(steps) * 130
    img, d = canvas(title, width, height)
    x1, x2 = 440, width - 440
    y = 130
    for idx, (text, color) in enumerate(steps):
        box(d, (x1, y, x2, y + 78), text, fill=color, fnt=BODY, wrap_width=34)
        if idx < len(steps) - 1:
            arrow(d, (width // 2, y + 82), (width // 2, y + 122))
        y += 130
    return save(img, name)


def auth_flow():
    steps = [
        ("User opens app", "green"), ("Sign in / Sign up", "blue"), ("Submit credentials", "blue"),
        ("Django Auth API validates", "orange"), ("JWT access + refresh token issued", "purple"),
        ("Frontend stores tokens", "gray"), ("Route user by role", "green"), ("Refresh token when needed", "orange"),
    ]
    return vertical_flow("Authentication and Role-Based Access Flow", steps, "05_authentication_role_flow")


def triage_flow():
    steps = [
        ("Emergency details captured", "green"), ("Triage API receives symptoms and location", "blue"),
        ("AI assigns P-level and care needs", "purple"), ("Search hospitals by distance, beds, ICU, services", "orange"),
        ("Rank recommended hospitals", "blue"), ("Notify dispatcher / reception", "green"),
        ("Create ambulance request if needed", "orange"), ("Track case and resolve", "gray"), ("Store outcome feedback", "purple"),
    ]
    return vertical_flow("Emergency AI Triage and Hospital Routing Flow", steps, "06_emergency_ai_triage_flow")


def sequence_diagram():
    img, d = canvas("Ambulance Booking and Dispatch Sequence", 2100, 1200)
    actors = ["Patient", "React Frontend", "Django API", "PostgreSQL", "Redis / WebSocket", "Driver App"]
    xs = [150, 520, 870, 1220, 1570, 1920]
    for x, actor in zip(xs, actors):
        box(d, (x - 120, 140, x + 120, 210), actor, fill="blue", fnt=SMALL, wrap_width=14)
        d.line((x, 220, x, 1080), fill="#B8C2CC", width=3)
    messages = [
        (0, 1, "Enter pickup and ambulance type"), (1, 2, "Create ambulance request"),
        (2, 3, "Save pending request"), (2, 4, "Publish notification"),
        (4, 5, "Notify nearby driver"), (5, 2, "Accept request"),
        (2, 3, "Assign ambulance"), (2, 4, "Broadcast status"),
        (4, 1, "Update tracking UI"), (5, 2, "Complete trip"), (2, 3, "Store duration"),
    ]
    y = 280
    for a, b, msg in messages:
        arrow(d, (xs[a], y), (xs[b], y))
        d.text(((xs[a] + xs[b]) / 2, y - 25), msg, fill=COLORS["ink"], font=TINY, anchor="mm")
        y += 70
    return save(img, "07_ambulance_booking_sequence")


def bed_flow():
    steps = [
        ("Search or register patient", "green"), ("Check suitable bed availability", "blue"),
        ("If no bed, create transfer request", "orange"), ("Select ward, department, and bed", "blue"),
        ("Create bed allocation", "purple"), ("Mark bed occupied", "red"), ("Monitor length of stay", "gray"),
        ("Create supervisor alert for long stay", "orange"), ("Discharge patient", "green"), ("Mark bed available", "blue"),
    ]
    return vertical_flow("Bed Admission and Discharge Flow", steps, "08_bed_admission_discharge_flow")


def transfer_flow():
    steps = [
        ("Reception creates transfer request", "green"), ("Enter patient, priority, bed and services", "blue"),
        ("Notify receiving hospital", "orange"), ("Review destination capacity", "purple"),
        ("Accept or reject transfer", "gray"), ("Book ambulance if required", "orange"),
        ("Move patient", "green"), ("Mark transfer completed", "blue"), ("Update admission / bed records", "purple"),
    ]
    return vertical_flow("Patient Transfer Request Flow", steps, "09_patient_transfer_request_flow")


def supervisor_flow():
    steps = [
        ("Scheduled or manual checks", "green"), ("Find long occupancy, resource mismatch, pending verification", "blue"),
        ("Create supervisor alerts", "orange"), ("Supervisor reviews dashboard", "purple"),
        ("Approve / reject hospital registration", "gray"), ("Resolve alerts or assign corrections", "green"),
        ("Store resolution history", "blue"),
    ]
    return vertical_flow("Supervisor Monitoring Flow", steps, "10_supervisor_monitoring_flow")


def realtime_flow():
    steps = [
        ("System event occurs", "green"), ("Django app publishes event", "blue"),
        ("Redis channel layer receives payload", "orange"), ("Django Channels pushes WebSocket message", "purple"),
        ("React dashboard updates cards, maps, and alerts", "green"),
    ]
    return vertical_flow("Real-Time Notification Flow", steps, "11_realtime_notification_flow")


def deployment():
    img, d = canvas("Deployment Diagram", 1800, 1200)
    nodes = [
        ("User Browser / Mobile", 700, 140, "green"),
        ("Nginx / HTTPS Reverse Proxy", 650, 300, "blue"),
        ("React Static Build", 210, 500, "gray"),
        ("Gunicorn Django REST API", 650, 500, "orange"),
        ("Daphne ASGI WebSocket", 1090, 500, "orange"),
        ("PostgreSQL", 300, 760, "purple"),
        ("Redis", 760, 760, "purple"),
        ("Celery Worker", 1180, 760, "blue"),
        ("Twilio / AI / Maps", 1180, 980, "green"),
    ]
    centers = {}
    for text, x, y, color in nodes:
        box(d, (x, y, x + 400, y + 95), text, fill=color, fnt=BODY, wrap_width=22)
        centers[text] = (x + 200, y + 48)
    def ar(a, b):
        arrow(d, centers[a], centers[b])
    ar("User Browser / Mobile", "Nginx / HTTPS Reverse Proxy")
    ar("Nginx / HTTPS Reverse Proxy", "React Static Build")
    ar("Nginx / HTTPS Reverse Proxy", "Gunicorn Django REST API")
    ar("Nginx / HTTPS Reverse Proxy", "Daphne ASGI WebSocket")
    ar("Gunicorn Django REST API", "PostgreSQL")
    ar("Gunicorn Django REST API", "Redis")
    ar("Daphne ASGI WebSocket", "Redis")
    ar("Redis", "Celery Worker")
    ar("Celery Worker", "Twilio / AI / Maps")
    return save(img, "12_deployment_diagram")


def build_docx(image_paths):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Project UML and Flow Diagrams")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Healthcare Management System")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph()
    for idx, path in enumerate(image_paths, start=1):
        heading = doc.add_paragraph()
        heading.style = "Heading 1"
        heading.add_run(path.stem.replace("_", " ").title())
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(7.1))
        if idx != len(image_paths):
            doc.add_page_break()

    doc.save(DOCX_OUT)


def main():
    builders = [
        architecture, use_case, er_diagram, class_diagram, auth_flow, triage_flow,
        sequence_diagram, bed_flow, transfer_flow, supervisor_flow, realtime_flow, deployment,
    ]
    paths = [builder() for builder in builders]
    build_docx(paths)
    print(DOCX_OUT)
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
