import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "database_schema_report.md"
DEFAULT_OUTPUT = ROOT / "docs" / "database_schema_report.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "D0D7DE"
CODE_FILL = "F6F8FA"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, font_size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = "Footer"
    run = footer.add_run("Database Schema Report")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)


def apply_document_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc, title_text, subtitle_text):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(title_text)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(subtitle_text)
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(90, 90, 90)


def clean_inline(text):
    return text.replace("**", "").replace("`", "")


def add_paragraph_with_inline_format(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            paragraph.add_run(part)
    return paragraph


def add_markdown_table(doc, table_lines):
    rows = []
    for line in table_lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [clean_inline(cell.strip()) for cell in stripped.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells):
            continue
        rows.append(cells)

    if not rows:
        return

    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    set_table_width(table)

    for r_idx, row in enumerate(rows):
        for c_idx in range(max_cols):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            set_cell_text(cell, text, bold=(r_idx == 0), font_size=8.5)
            set_cell_borders(cell)
            set_cell_shading(cell, LIGHT_BLUE if r_idx == 0 else "FFFFFF")

    doc.add_paragraph()


def add_code_block(doc, code_lines, language):
    label = "Database E-R Diagram Mermaid Source" if language == "mermaid" else "Code"
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(label)
    run.bold = True
    run.font.color.rgb = DARK_BLUE

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL)
    set_cell_borders(cell, "D8DEE4", "4")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(7.5 if language == "mermaid" else 8.5)
    run.font.color.rgb = RGBColor(36, 41, 47)
    doc.add_paragraph()


def convert_markdown_to_docx(source, output, title_text, subtitle_text):
    markdown = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    apply_document_styles(doc)
    add_title(doc, title_text, subtitle_text)

    table_buffer = []
    code_buffer = []
    in_code = False
    code_language = ""

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    for raw_line in markdown:
        line = raw_line.rstrip()

        if line.startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_language = line[3:].strip()
                code_buffer = []
            else:
                add_code_block(doc, code_buffer, code_language)
                in_code = False
                code_language = ""
                code_buffer = []
            continue

        if in_code:
            code_buffer.append(line)
            continue

        if line.strip().startswith("|"):
            table_buffer.append(line)
            continue
        flush_table()

        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            continue
        if stripped.startswith("## "):
            doc.add_heading(clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(clean_inline(stripped[4:]), level=2)
        elif stripped.startswith("- "):
            paragraph = add_paragraph_with_inline_format(doc, stripped[2:], style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.375)
            paragraph.paragraph_format.first_line_indent = Inches(-0.188)
            paragraph.paragraph_format.space_after = Pt(4)
        else:
            add_paragraph_with_inline_format(doc, stripped)

    flush_table()

    doc.save(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--title", default="Database Schema Report")
    parser.add_argument("--subtitle", default="Healthcare Management System")
    args = parser.parse_args()

    out = convert_markdown_to_docx(args.source, args.output, args.title, args.subtitle)
    print(out)
